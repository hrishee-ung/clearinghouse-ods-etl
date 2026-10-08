from __future__ import annotations

from pathlib import Path
from typing import Any, Dict

import pandas as pd
from openpyxl import load_workbook

from exceptions import ValidationError


def validate_source_data(
    data: dict[str, pd.DataFrame],
    config: Dict[str, Any],
) -> None:
    app_config = config.get("app", {})
    allow_zero_records = bool(app_config.get("allow_zero_records", False))

    if data is None:
        raise ValidationError("Source validation failed: extracted data is None.")

    if not isinstance(data, dict) or not data:
        raise ValidationError("Source validation failed: extracted data is empty or invalid.")

    total_records = sum(len(df.index) for df in data.values())
    if total_records == 0 and not allow_zero_records:
        raise ValidationError("Source validation failed: total extracted record count is zero.")

    required_columns = [
        "TERM_CODE",
        "UNG_ID",
        "SPRIDEN_FIRST_NAME",
        "NSC_MIDDLE_INITIAL",
        "SPRIDEN_LAST_NAME",
        "NAME_SUFFIX",
        "NSC_BIRTHDATE",
    ]

    for file_key, df in data.items():
        for col in required_columns:
            if col not in df.columns:
                raise ValidationError(
                    f"Required column '{col}' missing from extract for file '{file_key}'."
                )


def validate_clearinghouse_files(
    file_artifacts: list[dict[str, Any]],
    config: Dict[str, Any],
) -> None:
    validation_config = config.get("validation", {})
    file_must_exist = bool(validation_config.get("file_must_exist", True))
    file_must_not_be_empty = bool(validation_config.get("file_must_not_be_empty", True))

    if not file_artifacts:
        raise ValidationError("No generated file artifacts found for validation.")

    for artifact in file_artifacts:
        workbook_path = Path(artifact["workbook_path"])
        text_file_path = Path(artifact["text_file_path"])

        if file_must_exist and not workbook_path.exists():
            raise ValidationError(f"Generated workbook does not exist: {workbook_path}")

        if file_must_exist and not text_file_path.exists():
            raise ValidationError(f"Generated text file does not exist: {text_file_path}")

        if file_must_not_be_empty and workbook_path.stat().st_size == 0:
            raise ValidationError(f"Generated workbook is empty: {workbook_path}")

        if file_must_not_be_empty and text_file_path.stat().st_size == 0:
            raise ValidationError(f"Generated text file is empty: {text_file_path}")

        _validate_workbook(workbook_path)
        _validate_text_file(text_file_path, artifact["record_count"])


def _validate_workbook(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    if ws["A1"].value != "H1":
        raise ValidationError(f"Workbook header row in '{workbook_path.name}' does not begin with H1.")

    if ws[f"A{ws.max_row}"].value != "T1":
        raise ValidationError(f"Workbook trailer row in '{workbook_path.name}' does not begin with T1.")

    expected_total_rows = ws.max_row
    trailer_count_value = ws[f"B{ws.max_row}"].value

    if str(trailer_count_value) != str(expected_total_rows):
        raise ValidationError(
            f"Workbook trailer count mismatch in '{workbook_path.name}'. "
            f"Expected {expected_total_rows}, found {trailer_count_value}."
        )


def _validate_text_file(text_file_path: Path, expected_data_rows: int) -> None:
    with text_file_path.open("r", encoding="utf-8", newline="") as f:
        lines = f.readlines()

    if not lines:
        raise ValidationError(f"Generated text file is empty: {text_file_path}")

    if not lines[0].startswith("H1\t") and not lines[0].startswith("H1"):
        raise ValidationError(f"Text file header row invalid in '{text_file_path.name}'.")

    if not lines[-1].startswith("T1\t") and not lines[-1].startswith("T1"):
        raise ValidationError(f"Text file trailer row invalid in '{text_file_path.name}'.")

    expected_total_lines = expected_data_rows + 2
    actual_total_lines = len(lines)

    if actual_total_lines != expected_total_lines:
        raise ValidationError(
            f"Text file line count mismatch in '{text_file_path.name}'. "
            f"Expected {expected_total_lines}, found {actual_total_lines}."
        )